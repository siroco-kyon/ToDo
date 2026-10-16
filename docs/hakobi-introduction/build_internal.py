"""Build the HAKOBI internal report (Word) from report-internal.md.

The PDF is made from the generated Word file (LibreOffice or Word).
Markdown extensions used by the source:
  > text              cover meta line (only right after the title)
  :::summary TITLE    shaded summary box, closed by :::
  ## 付録A 名前        appendix chapter (labelled, no number)
"""

from __future__ import annotations

import argparse
import hashlib
import json
from dataclasses import dataclass
from pathlib import Path
import re

from PIL import Image
import pypdfium2 as pdfium
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parent
REPO = ROOT.parent.parent
ASSETS = ROOT / "assets-current"
BUILD = ROOT / ".build" / "internal"
ICON = REPO / "build" / "icon.png"

FONT = "Yu Gothic"
TEXT_WIDTH_MM = 174.0
BODY_PT = 10.5
BODY_LINE_PT = 17.5
FOOTER_TEXT = "HAKOBI 社内レポート"

# Colors follow the app icon: navy ground, teal bars and an amber accent.
NAVY = "0F1B2D"
INK = "1E2B37"
MUTED = "5B6B78"
TEAL = "0E8C80"
TEAL_DARK = "0B6E66"
TEAL_LIGHT = "E7F5F3"
AMBER = "E9A21B"
AMBER_LIGHT = "FFF6E0"
SLATE = "6B7C86"
SLATE_LIGHT = "F1F4F6"
RULE = "D5DCE2"

# Paragraphs whose last line would hold only "す。". Tighten their letter spacing
# slightly (twentieths of a point); this changes typography, not the text.
TIGHTEN = {
    "HAKOBIはこうした手間": -8,
    "すべての画面を覚える必要はなく": -8,
    "チームで共有する機能のほかに": -8,
    "作業の区切りに残した短い記録を": -8,
}


def tighten(paragraph, text: str) -> None:
    for prefix, value in TIGHTEN.items():
        if text.startswith(prefix):
            for run in paragraph.runs:
                element = OxmlElement("w:spacing")
                element.set(qn("w:val"), str(value))
                properties = run._r.get_or_add_rPr()
                # OOXML order: color, spacing, ..., sz.
                size = properties.find(qn("w:sz"))
                if size is None:
                    properties.append(element)
                else:
                    size.addprevious(element)


# Chapters that start on a new page. The others continue after the previous chapter.
NEW_PAGE_CHAPTERS = {
    "作ったきっかけ", "チームで使うと変わること", "記録を定例と週報に使う", "自分の仕事も整理しやすい",
    "仕組みとデータの扱い", "導入の進め方", "付録A 週報用エージェントの作り方",
}


@dataclass(frozen=True)
class ImageSpec:
    # Fractions of the source image, in left/top/right/bottom order.
    crop: tuple[float, float, float, float] | None = None
    width_mm: float = 168.0
    max_height_mm: float = 110
    diagram: bool = False


# Screenshot sources are 1759 x 1244. Crop to the part each section explains.
IMAGE_SPECS: dict[str, ImageSpec] = {
    "gantt.png": ImageSpec((0.0, 0.0450161, 0.704946, 0.5916399)),
    "team.png": ImageSpec((0.0085276, 0.0956592, 0.5025583, 0.4863344), width_mm=136),
    "progress.png": ImageSpec((0.2615122, 0.1366559, 0.7305287, 0.5747588), width_mm=122),
    "report.png": ImageSpec((0.1051734, 0.0522508, 0.8857305, 0.4823151)),
    "meeting.png": ImageSpec((0.0, 0.0, 0.9977260, 0.3054662)),
    "plan.png": ImageSpec((0.0085276, 0.0602894, 0.6139852, 0.4340836), width_mm=140),
    "record-hub.png": ImageSpec(width_mm=164, diagram=True),
    "weekly-flow-v2.png": ImageSpec(width_mm=164, diagram=True),
    "architecture-v2.png": ImageSpec(width_mm=150, diagram=True),
    "migration-flow-v2.png": ImageSpec(width_mm=164, diagram=True),
}


# ---------------------------------------------------------------- OOXML helpers

def xml_bool(parent, name: str, value: bool) -> None:
    for element in list(parent.findall(qn("w:" + name))):
        parent.remove(element)
    element = OxmlElement("w:" + name)
    element.set(qn("w:val"), "1" if value else "0")
    parent.append(element)


def set_run_font(properties, *, size: float, bold: bool = False, color: str = INK) -> None:
    rfonts = properties.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        properties.insert(0, rfonts)
    for name in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn("w:" + name), FONT)
    for name in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rfonts.attrib.pop(qn("w:" + name), None)
    for name in ("b", "bCs", "color", "sz", "szCs"):
        for old in list(properties.findall(qn("w:" + name))):
            properties.remove(old)
    xml_bool(properties, "b", bold)
    xml_bool(properties, "bCs", bold)
    color_element = OxmlElement("w:color")
    color_element.set(qn("w:val"), color)
    properties.append(color_element)
    for name in ("sz", "szCs"):
        element = OxmlElement("w:" + name)
        element.set(qn("w:val"), str(round(size * 2)))
        properties.append(element)


def style_font(style, size: float, *, bold: bool = False, color: str = INK) -> None:
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor.from_string(color)
    set_run_font(style.element.get_or_add_rPr(), size=size, bold=bold, color=color)


def no_grid(paragraph) -> None:
    xml_bool(paragraph._p.get_or_add_pPr(), "snapToGrid", False)


def paragraph_border(paragraph, side: str, *, color: str, size_eighths: int, space: int = 4) -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    element = OxmlElement("w:" + side)
    for key, value in (("val", "single"), ("sz", str(size_eighths)), ("space", str(space)), ("color", color)):
        element.set(qn("w:" + key), value)
    borders.append(element)


def spacing(paragraph, *, before: float = 0, after: float = 0, line: float | None = None) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    if line is not None:
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        fmt.line_spacing = Pt(line)


def add_run(paragraph, text: str, *, size: float = BODY_PT, bold: bool = False, color: str = INK):
    run = paragraph.add_run(text)
    set_run_font(run._r.get_or_add_rPr(), size=size, bold=bold, color=color)
    return run


INLINE_PATTERN = re.compile(r"(\*\*.*?\*\*|`[^`]+`|\[[^\]]+\]\([^\)]+\))")


def inline_text(paragraph, text: str, *, size: float = BODY_PT, color: str = INK, bold: bool = False) -> None:
    for piece in INLINE_PATTERN.split(text):
        if not piece:
            continue
        link = re.fullmatch(r"\[([^\]]+)\]\(([^\)]+)\)", piece)
        if link:
            node = OxmlElement("w:hyperlink")
            node.set(qn("r:id"), paragraph.part.relate_to(link.group(2), RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
            run_node = OxmlElement("w:r")
            properties = OxmlElement("w:rPr")
            set_run_font(properties, size=size, color=TEAL_DARK)
            underline = OxmlElement("w:u")
            underline.set(qn("w:val"), "single")
            properties.append(underline)
            run_node.append(properties)
            value = OxmlElement("w:t")
            value.text = link.group(1)
            run_node.append(value)
            node.append(run_node)
            paragraph._p.append(node)
            continue
        is_bold = piece.startswith("**") and piece.endswith("**")
        if is_bold:
            piece = piece[2:-2]
        elif piece.startswith("`") and piece.endswith("`"):
            piece = piece[1:-1]
        add_run(paragraph, piece, size=size, bold=bold or is_bold, color=color)


def cell_fill(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    for old in list(properties.findall(qn("w:shd"))):
        properties.remove(old)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def cell_margins(cell, *, top: int, bottom: int, left: int, right: int) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for name, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        element = OxmlElement("w:" + name)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    properties.append(margins)


def cell_borders(cell, **sides: tuple[int, str] | None) -> None:
    """sides: top/left/bottom/right = (size in eighths of a point, color) or None for no line."""
    properties = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        element = OxmlElement("w:" + side)
        value = sides.get(side)
        if value is None:
            element.set(qn("w:val"), "nil")
        else:
            size, color = value
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), str(size))
            element.set(qn("w:color"), color)
            element.set(qn("w:space"), "0")
        borders.append(element)
    properties.append(borders)


def table_borders(table, color: str | None) -> None:
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement("w:" + name)
        if color is None:
            element.set(qn("w:val"), "nil")
        else:
            for key, value in (("val", "single"), ("sz", "4"), ("color", color)):
                element.set(qn("w:" + key), value)
        borders.append(element)
    properties.append(borders)


def fixed_table(doc, rows: int, widths_mm: list[float]):
    table = doc.add_table(rows=rows, cols=len(widths_mm))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)
    for column, width in zip(table.columns, widths_mm):
        column.width = Mm(width)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row.cells, widths_mm):
            cell.width = Mm(width)
    return table


def cell_paragraph(cell, *, first: bool):
    paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
    no_grid(paragraph)
    return paragraph


def gap(doc, points: float) -> None:
    paragraph = doc.add_paragraph()
    no_grid(paragraph)
    spacing(paragraph, line=max(points, 1))
    add_run(paragraph, "", size=1)


# ---------------------------------------------------------------- document setup

def prepare_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin, section.bottom_margin = Mm(17), Mm(17)
    section.left_margin = section.right_margin = Mm(18)
    section.header_distance, section.footer_distance = Mm(6), Mm(8)
    for grid in list(section._sectPr.findall(qn("w:docGrid"))):
        section._sectPr.remove(grid)
    for style in doc.styles:
        if style.type == 1:
            xml_bool(style.element.get_or_add_pPr(), "snapToGrid", False)

    normal = doc.styles["Normal"]
    style_font(normal, BODY_PT)
    fmt = normal.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(BODY_LINE_PT)
    fmt.space_before, fmt.space_after = Pt(0), Pt(7)
    fmt.widow_control = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style_font(style, BODY_PT)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        style.paragraph_format.line_spacing = Pt(BODY_LINE_PT)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(3)

    doc.core_properties.title = "チームの仕事が見える業務管理ツール HAKOBI"
    doc.core_properties.subject = "HAKOBIを作った目的、チームで使うと変わること、導入の進め方"
    doc.core_properties.keywords = "HAKOBI ガント 進捗 見える化 定例 週報 Copilot"
    doc.core_properties.author = ""
    doc.core_properties.comments = ""

    footer = section.footer.paragraphs[0]
    no_grid(footer)
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    spacing(footer, line=11)
    add_run(footer, FOOTER_TEXT + "    ", size=8, color=MUTED)
    for index, instruction in enumerate(("PAGE", "NUMPAGES")):
        if index:
            add_run(footer, " / ", size=8, color=MUTED)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), instruction)
        run = OxmlElement("w:r")
        properties = OxmlElement("w:rPr")
        set_run_font(properties, size=8, color=MUTED)
        run.append(properties)
        value = OxmlElement("w:t")
        value.text = "1"
        run.append(value)
        field.append(run)
        footer._p.append(field)
    return doc


def new_numbering(doc, *, decimal: bool) -> int:
    root = doc.part.numbering_part.element
    abstract_ids = [int(e.get(qn("w:abstractNumId"))) for e in root.findall(qn("w:abstractNum"))]
    num_ids = [int(e.get(qn("w:numId"))) for e in root.findall(qn("w:num"))]
    abstract_id, num_id = max(abstract_ids, default=-1) + 1, max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multilevel = OxmlElement("w:multiLevelType")
    multilevel.set(qn("w:val"), "singleLevel")
    abstract.append(multilevel)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for name, value in (("start", "1"), ("numFmt", "decimal" if decimal else "bullet"),
                        ("lvlText", "%1." if decimal else "●"), ("lvlJc", "left")):
        element = OxmlElement("w:" + name)
        element.set(qn("w:val"), value)
        level.append(element)
    properties = OxmlElement("w:pPr")
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "340")
    indent.set(qn("w:hanging"), "280")
    properties.append(indent)
    level.append(properties)
    run_properties = OxmlElement("w:rPr")
    set_run_font(run_properties, size=BODY_PT if decimal else 7, bold=decimal, color=TEAL)
    level.append(run_properties)
    abstract.append(level)
    first_number = root.find(qn("w:num"))
    if first_number is None:
        root.append(abstract)
    else:
        root.insert(list(root).index(first_number), abstract)
    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(num_id))
    reference = OxmlElement("w:abstractNumId")
    reference.set(qn("w:val"), str(abstract_id))
    number.append(reference)
    root.append(number)
    return num_id


def list_item(doc, text: str, *, number_id: int, decimal: bool) -> None:
    paragraph = doc.add_paragraph(style="List Number" if decimal else "List Bullet")
    no_grid(paragraph)
    properties = paragraph._p.get_or_add_pPr()
    for old in list(properties.findall(qn("w:numPr"))):
        properties.remove(old)
    numbering = OxmlElement("w:numPr")
    for name, value in (("ilvl", "0"), ("numId", str(number_id))):
        element = OxmlElement("w:" + name)
        element.set(qn("w:val"), value)
        numbering.append(element)
    properties.append(numbering)
    inline_text(paragraph, text)


# ---------------------------------------------------------------- blocks

def add_cover(doc, title: str, meta: str) -> None:
    table = fixed_table(doc, 1, [TEXT_WIDTH_MM])
    table_borders(table, None)
    cell = table.cell(0, 0)
    cell_fill(cell, NAVY)
    cell_margins(cell, top=420, bottom=440, left=480, right=480)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    icon_paragraph = cell_paragraph(cell, first=True)
    spacing(icon_paragraph, after=10)
    icon_paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    icon_paragraph.add_run().add_picture(str(ICON), width=Mm(15))

    meta_paragraph = cell_paragraph(cell, first=False)
    spacing(meta_paragraph, after=6, line=14)
    add_run(meta_paragraph, meta, size=9.5, bold=True, color="7FE3D3")

    # Break the title after the tool category, so the product name stands on its own line.
    head, _, tail = title.partition("業務管理ツール")
    lines = [head + "業務管理ツール", tail.strip()] if tail else [title]
    for index, line in enumerate(lines):
        paragraph = cell_paragraph(cell, first=False)
        spacing(paragraph, after=0, line=34)
        add_run(paragraph, line, size=24, bold=True, color="FFFFFF")

    gap(doc, 10)


def add_box(doc, title: str, lines: list[str], *, fill: str, bar: str, title_color: str) -> None:
    table = fixed_table(doc, 1, [TEXT_WIDTH_MM])
    table_borders(table, None)
    cell = table.cell(0, 0)
    cell_fill(cell, fill)
    cell_borders(cell, left=(36, bar), top=None, bottom=None, right=None)
    cell_margins(cell, top=200, bottom=200, left=300, right=300)
    heading = cell_paragraph(cell, first=True)
    spacing(heading, after=4, line=17)
    add_run(heading, title, size=11, bold=True, color=title_color)
    for line in lines:
        numbered = re.match(r"^(\d+)\.\s+(.+)$", line)
        paragraph = cell_paragraph(cell, first=False)
        spacing(paragraph, after=3, line=17)
        if numbered:
            paragraph.paragraph_format.left_indent = Mm(6)
            paragraph.paragraph_format.first_line_indent = Mm(-6)
            tab_stops = paragraph.paragraph_format.tab_stops
            tab_stops.add_tab_stop(Mm(6))
            add_run(paragraph, numbered.group(1) + "\t", bold=True, color=title_color)
            inline_text(paragraph, numbered.group(2))
            tighten(paragraph, numbered.group(2))
        else:
            inline_text(paragraph, line)
    gap(doc, 8)


def add_chapter(doc, heading: str, number: int | None, *, page_break: bool) -> None:
    label = None
    title = heading
    appendix = re.match(r"^(付録[A-Z])\s+(.+)$", heading)
    if appendix:
        label, title = appendix.group(1), appendix.group(2)
    elif number is not None:
        label = f"{number:02d}"
    paragraph = doc.add_paragraph()
    no_grid(paragraph)
    paragraph.paragraph_format.page_break_before = page_break
    paragraph.paragraph_format.keep_with_next = True
    spacing(paragraph, before=0 if page_break else 24, after=12, line=28)
    if label:
        add_run(paragraph, label, size=15, bold=True, color=TEAL)
        add_run(paragraph, "　", size=19, color=NAVY)
    add_run(paragraph, title, size=19, bold=True, color=NAVY)
    paragraph_border(paragraph, "bottom", color=TEAL, size_eighths=12, space=5)


def add_section(doc, text: str) -> None:
    paragraph = doc.add_paragraph()
    no_grid(paragraph)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.left_indent = Mm(2.5)
    spacing(paragraph, before=9, after=5, line=19)
    paragraph_border(paragraph, "left", color=TEAL, size_eighths=24, space=6)
    add_run(paragraph, text, size=12.5, bold=True, color=NAVY)


def add_paragraph(doc, text: str, *, size: float = BODY_PT, color: str = INK):
    paragraph = doc.add_paragraph()
    no_grid(paragraph)
    inline_text(paragraph, text, size=size, color=color)
    tighten(paragraph, text)
    return paragraph


TABLE_RATIOS = {
    "場面": (0.22, 0.25, 0.53),
    "タイミング": (0.17, 0.56, 0.27),
    "機能": (0.24, 0.76),
    "比較する項目": (0.22, 0.33, 0.45),
    "書く場所": (0.15, 0.30, 0.55),
    "部分": (0.24, 0.76),
}

HEADER_FILLS = {"これまで": SLATE, "HAKOBIを使うと": TEAL}


def add_table(doc, rows: list[list[str]]) -> None:
    columns = len(rows[0])
    if any(len(row) != columns for row in rows):
        raise ValueError("Markdown table rows have inconsistent cell counts")
    ratios = TABLE_RATIOS.get(rows[0][0], tuple(1 / columns for _ in range(columns)))
    widths = [TEXT_WIDTH_MM * ratio for ratio in ratios]
    table = fixed_table(doc, len(rows), widths)
    table_borders(table, RULE)
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        if row_index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for column_index, value in enumerate(values):
            cell = row.cells[column_index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            header = row_index == 0
            if header:
                fill = HEADER_FILLS.get(value, NAVY)
            elif column_index == 0:
                fill = SLATE_LIGHT
            else:
                fill = "FFFFFF"
            cell_fill(cell, fill)
            cell_margins(cell, top=100, bottom=100, left=120, right=120)
            paragraph = cell_paragraph(cell, first=True)
            spacing(paragraph, line=15)
            paragraph.paragraph_format.keep_with_next = header
            inline_text(paragraph, value, size=9.5,
                        color="FFFFFF" if header else (NAVY if column_index == 0 else INK),
                        bold=header or column_index == 0)
    gap(doc, 8)


def add_code(doc, lines: list[str]) -> None:
    table = fixed_table(doc, 1, [TEXT_WIDTH_MM])
    table_borders(table, None)
    cell = table.cell(0, 0)
    cell_fill(cell, "F4F6F8")
    cell_borders(cell, top=(4, RULE), bottom=(4, RULE), left=(4, RULE), right=(4, RULE))
    cell_margins(cell, top=160, bottom=160, left=240, right=240)
    # Long instruction text may continue on the next page; short blocks stay together.
    if len(lines) > 12:
        for tr in table._tbl.findall(qn("w:tr")):
            properties = tr.get_or_add_trPr()
            for old in list(properties.findall(qn("w:cantSplit"))):
                properties.remove(old)
    for index, line in enumerate(lines):
        paragraph = cell_paragraph(cell, first=index == 0)
        spacing(paragraph, line=14.5)
        add_run(paragraph, line, size=9.5, color=INK)
    gap(doc, 8)


# ---------------------------------------------------------------- figures

def register_fonts() -> None:
    pdfmetrics.registerFont(TTFont("YuR", "C:/Windows/Fonts/YuGothR.ttc", subfontIndex=0))
    pdfmetrics.registerFont(TTFont("YuB", "C:/Windows/Fonts/YuGothB.ttc", subfontIndex=0))


def hexc(value: str):
    return colors.HexColor("#" + value)


def text(c, x, y, value, *, size=10, bold=False, color=INK, align="center"):
    c.setFillColor(hexc(color))
    c.setFont("YuB" if bold else "YuR", size)
    {"center": c.drawCentredString, "left": c.drawString, "right": c.drawRightString}[align](x, y, value)


def node(c, x, y, w, h, title, detail=(), *, fill="FFFFFF", stroke=TEAL, title_color=NAVY, detail_color=MUTED, title_size=11):
    c.setFillColor(hexc(fill))
    c.setStrokeColor(hexc(stroke))
    c.setLineWidth(1)
    c.roundRect(x, y, w, h, 6, stroke=1, fill=1)
    if detail:
        title_y = y + h / 2 + 3
        text(c, x + w / 2, title_y, title, size=title_size, bold=True, color=title_color)
        for index, line in enumerate(detail):
            text(c, x + w / 2, title_y - 16 - index * 13, line, size=8.6, color=detail_color)
    else:
        text(c, x + w / 2, y + h / 2 - 4, title, size=title_size, bold=True, color=title_color)


def arrow(c, x1, y1, x2, y2, *, color=SLATE, double=False):
    import math
    c.setStrokeColor(hexc(color))
    c.setFillColor(hexc(color))
    c.setLineWidth(1.2)
    c.line(x1, y1, x2, y2)
    angle = math.atan2(y2 - y1, x2 - x1)

    def head(x, y, a):
        path = c.beginPath()
        path.moveTo(x, y)
        path.lineTo(x - 6 * math.cos(a - .42), y - 6 * math.sin(a - .42))
        path.lineTo(x - 6 * math.cos(a + .42), y - 6 * math.sin(a + .42))
        path.close()
        c.drawPath(path, fill=1, stroke=0)

    head(x2, y2, angle)
    if double:
        head(x1, y1, angle + math.pi)


def draw_record_hub(c) -> None:
    text(c, 92, 186, "作業の区切りに1回残す", size=9.5, bold=True, color=TEAL_DARK)
    node(c, 10, 62, 164, 104, "日々の記録", ("進捗率・進捗ログ", "作業時間"),
         fill=TEAL, stroke=TEAL, title_color="FFFFFF", detail_color="E6FFFB", title_size=14)
    text(c, 380, 186, "チームの確認と報告に何度も使う", size=9.5, bold=True, color=TEAL_DARK)
    outputs = (
        ("チームの状況確認", "ガント・チーム画面"),
        ("定例での相談", "報告タブ・定例モード"),
        ("週報の下書き", "報告の文章をCopilotへ"),
    )
    for index, (title, detail) in enumerate(outputs):
        y = 124 - index * 56
        node(c, 248, y, 256, 46, title, (detail,), fill=TEAL_LIGHT, stroke=TEAL)
        arrow(c, 176, 114, 245, y + 23, color=TEAL)


def draw_weekly(c) -> None:
    steps = (
        ("報告タブ", "期間と担当を選ぶ", TEAL_LIGHT, TEAL),
        ("Markdownをコピー", "見出しつきの文章", TEAL_LIGHT, TEAL),
        ("週報用エージェント", "記録から下書き", "EEF2F7", SLATE),
        ("本人が確認", "整えて提出", AMBER_LIGHT, AMBER),
    )
    for index, (title, detail, fill, stroke) in enumerate(steps):
        x = 6 + index * 127
        node(c, x, 22, 112, 62, title, (detail,), fill=fill, stroke=stroke, title_size=10.5)
        if index < 3:
            arrow(c, x + 114, 53, x + 125, 53)
    for x1, x2, label, color in ((6, 245, "HAKOBI", TEAL_DARK), (260, 372, "Microsoft 365 Copilot", SLATE),
                                 (387, 499, "担当者", "9A6A0B")):
        c.setStrokeColor(hexc(color))
        c.setLineWidth(1)
        c.line(x1, 96, x2, 96)
        c.line(x1, 96, x1, 91)
        c.line(x2, 96, x2, 91)
        text(c, (x1 + x2) / 2, 103, label, size=9, bold=True, color=color)


def draw_architecture(c) -> None:
    text(c, 18, 190, "個人用", size=12, bold=True, color=NAVY, align="left")
    text(c, 214, 190, "チーム用", size=12, bold=True, color=NAVY, align="left")
    c.setStrokeColor(hexc(RULE))
    c.setLineWidth(1)
    c.line(194, 12, 194, 196)
    node(c, 16, 74, 160, 86, "自分のPC", ("デスクトップ版", "PC内にデータを保存"), fill=SLATE_LIGHT, stroke=SLATE)
    text(c, 96, 46, "1人で使う", size=10, color=MUTED)
    node(c, 214, 124, 128, 42, "ブラウザ", fill="FFFFFF", stroke=SLATE)
    node(c, 368, 124, 128, 42, "デスクトップ版", fill="FFFFFF", stroke=SLATE)
    node(c, 262, 14, 186, 64, "社内LANのサーバー", ("タスク・進捗・作業記録",), fill=TEAL, stroke=TEAL,
         title_color="FFFFFF", detail_color="E6FFFB")
    arrow(c, 280, 122, 318, 80, double=True)
    arrow(c, 430, 122, 392, 80, double=True)
    text(c, 355, 100, "同じデータを使う", size=8.6, color=MUTED)


def draw_migration(c) -> None:
    steps = (
        ("個人用データ", "todo.dbを用意", SLATE_LIGHT, SLATE),
        ("管理者が確認", "取り込み先・ドライラン", TEAL_LIGHT, TEAL),
        ("チーム用へ取り込み", "以後はチームで使う", TEAL_LIGHT, TEAL),
    )
    for index, (title, detail, fill, stroke) in enumerate(steps):
        x = 6 + index * 170
        node(c, x, 30, 150, 58, title, (detail,), fill=fill, stroke=stroke)
        if index < 2:
            arrow(c, x + 153, 59, x + 167, 59)
    text(c, 256, 9, "個人用DBとチーム用DBは自動で同期しません", size=9, color=MUTED)


FIGURES = (
    ("record-hub.png", 512, 200, draw_record_hub),
    ("weekly-flow-v2.png", 512, 116, draw_weekly),
    ("architecture-v2.png", 512, 204, draw_architecture),
    ("migration-flow-v2.png", 512, 98, draw_migration),
)


def generate_figures() -> None:
    register_fonts()
    directory = BUILD / "figures"
    directory.mkdir(parents=True, exist_ok=True)
    for name, width, height, draw in FIGURES:
        pdf_path = directory / (Path(name).stem + ".pdf")
        c = canvas.Canvas(str(pdf_path), pagesize=(width, height))
        c.setFillColor(colors.white)
        c.rect(0, 0, width, height, fill=1, stroke=0)
        draw(c)
        c.save()
        pdf = pdfium.PdfDocument(str(pdf_path))
        page = pdf[0]
        bitmap = page.render(scale=3)
        bitmap.to_pil().save(ASSETS / name)
        bitmap.close()
        page.close()
        pdf.close()


def prepare_image(reference: str, manifest: list[dict]) -> tuple[Path, float]:
    source = ROOT / reference
    if not source.is_file():
        raise FileNotFoundError(f"Required image is missing: {source}")
    spec = IMAGE_SPECS.get(source.name, ImageSpec())
    with Image.open(source) as original:
        original.load()
        width, height = original.size
        left, top, right, bottom = 0, 0, width, height
        if spec.crop:
            a, b, c, d = spec.crop
            left, top, right, bottom = round(width * a), round(height * b), round(width * c), round(height * d)
        max_pixels_high = round((right - left) * spec.max_height_mm / spec.width_mm)
        if bottom - top > max_pixels_high:
            if spec.diagram:
                raise ValueError(f"Diagram {source.name} is taller than {spec.max_height_mm} mm")
            bottom = top + max_pixels_high
        cropped = original.crop((left, top, right, bottom))
        destination = BUILD / "crops" / source.name
        destination.parent.mkdir(parents=True, exist_ok=True)
        cropped.save(destination)
        displayed_height = spec.width_mm * cropped.height / cropped.width
    manifest.append({
        "reference": reference,
        "crop_pixels": [left, top, right, bottom],
        "display_mm": [spec.width_mm, round(displayed_height, 1)],
    })
    return destination, spec.width_mm


def add_image(doc, reference: str, caption: str, manifest: list[dict]) -> None:
    path, width_mm = prepare_image(reference, manifest)
    paragraph = doc.add_paragraph()
    no_grid(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    spacing(paragraph, before=3, after=3)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.paragraph_format.keep_with_next = True
    picture = paragraph.add_run().add_picture(str(path), width=Mm(width_mm))
    picture._inline.docPr.set("descr", caption)
    caption_paragraph = add_paragraph(doc, caption, size=8.5, color=MUTED)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    spacing(caption_paragraph, after=10, line=12.5)


# ---------------------------------------------------------------- markdown

def table_separator(cells: list[str]) -> bool:
    return all(re.fullmatch(r":?-+:?", value.replace(" ", "")) for value in cells)


def build_document(source: Path, output: Path) -> dict:
    raw = source.read_bytes()
    lines = raw.decode("utf-8-sig").splitlines()
    BUILD.mkdir(parents=True, exist_ok=True)
    generate_figures()
    doc = prepare_document()
    manifest: list[dict] = []
    index = 0
    chapter_number = 0
    title = None
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("# "):
            title = line[2:].strip()
            index += 1
            meta = ""
            while index < len(lines) and not lines[index].strip():
                index += 1
            if index < len(lines) and lines[index].strip().startswith("> "):
                meta = lines[index].strip()[2:]
                index += 1
            add_cover(doc, title, meta)
            continue
        if line.startswith(":::"):
            kind, _, box_title = line[3:].partition(" ")
            index += 1
            body = []
            while index < len(lines) and lines[index].strip() != ":::":
                if lines[index].strip():
                    body.append(lines[index].strip())
                index += 1
            index += 1
            styles = {
                "summary": (TEAL_LIGHT, TEAL, TEAL_DARK),
                "note": (AMBER_LIGHT, AMBER, "8A5A00"),
            }
            fill, bar, color = styles.get(kind, styles["summary"])
            add_box(doc, box_title.strip(), body, fill=fill, bar=bar, title_color=color)
            continue
        if line.startswith("## "):
            heading = line[3:].strip()
            numbered = not heading.startswith("付録") and heading != "おわりに"
            if numbered:
                chapter_number += 1
            add_chapter(doc, heading, chapter_number if numbered else None,
                        page_break=heading in NEW_PAGE_CHAPTERS)
            index += 1
            continue
        if line.startswith("### "):
            add_section(doc, line[4:].strip())
            index += 1
            continue
        if line.startswith("```"):
            index += 1
            block = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            index += 1
            add_code(doc, block)
            continue
        if line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [value.strip() for value in lines[index].strip().strip("|").split("|")]
                if not table_separator(cells):
                    rows.append(cells)
                index += 1
            add_table(doc, rows)
            continue
        image = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", line)
        if image:
            add_image(doc, image.group(2), image.group(1), manifest)
            index += 1
            continue
        decimal = re.match(r"^\d+\.\s+(.+)$", line)
        bullet = re.match(r"^-\s+(.+)$", line)
        if decimal or bullet:
            is_decimal = bool(decimal)
            number_id = new_numbering(doc, decimal=is_decimal)
            items = []
            while index < len(lines):
                item = re.match(r"^\d+\.\s+(.+)$" if is_decimal else r"^-\s+(.+)$", lines[index].strip())
                if not item:
                    break
                items.append(item.group(1))
                index += 1
            for position, item in enumerate(items):
                list_item(doc, item, number_id=number_id, decimal=is_decimal)
            doc.paragraphs[-1].paragraph_format.space_after = Pt(8)
            continue
        paragraph = add_paragraph(doc, line)
        if re.fullmatch(r"\[[^\]]+\]\([^\)]+\)", line):
            spacing(paragraph, after=2)
        index += 1
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    result = {
        "docx": str(output),
        "source": str(source),
        "source_sha256": hashlib.sha256(raw).hexdigest(),
        "chapters": chapter_number,
        "images": manifest,
    }
    (BUILD / "build-manifest.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=ROOT / "report-internal.md")
    parser.add_argument("--output", type=Path, default=ROOT / "HAKOBI_社内レポート.docx")
    arguments = parser.parse_args()
    source = arguments.source if arguments.source.is_absolute() else ROOT / arguments.source
    output = arguments.output if arguments.output.is_absolute() else ROOT / arguments.output
    result = build_document(source, output)
    print(json.dumps({key: value for key, value in result.items() if key != "images"}, ensure_ascii=False))


if __name__ == "__main__":
    main()
